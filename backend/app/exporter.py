from __future__ import annotations

import csv
import io
import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from .payroll.schemas import PayrollPage
from .validation.payroll import validate_payroll_pages
from .validation.timecard import validate_timecard_transcription


HEADER_FILL = PatternFill("solid", fgColor="173772")
WARNING_FILL = PatternFill("solid", fgColor="FFF3CD")
ERROR_FILL = PatternFill("solid", fgColor="F8D7DA")
ERROR_BORDER = Border(
    left=Side(
        style="medium",
        color="DC3545",
    )
)


def _rows(value: dict[str, Any], document_type: str) -> tuple[list[str], list[list[str]]]:
    if document_type == "cartao-ponto":
        days = [day for page in value.get("pages", []) for day in page.get("days", [])]
        max_punches = max((len(day.get("punches", [])) for day in days), default=0)
        headers = ["Data"]
        for index in range(max_punches):
            headers.append(("Entrada " if index % 2 == 0 else "Saída ") + str(index // 2 + 1))
        rows = []
        for day in days:
            times = [punch.get("time_hhmm", "") for punch in day.get("punches", [])]
            rows.append([day.get("date_raw", ""), *times, *([""] * (max_punches - len(times)))])
        return headers, rows

    pages = value.get("pages", [])
    labels: list[str] = []
    for page in pages:
        for field in page.get("fields", []):
            label = field.get("label", "")
            if label and label not in labels:
                labels.append(label)
    headers = ["Pág.", "Mês", "Ano", *labels]
    rows = []
    for page in pages:
        by_label = {field.get("label", ""): field.get("value", "") for field in page.get("fields", [])}
        rows.append([str(page.get("page", "")), page.get("month", ""), page.get("year", ""),
                     *[by_label.get(label, "") for label in labels]])
    return headers, rows


def export_json(value: dict[str, Any]) -> bytes:
    return json.dumps(value, ensure_ascii=False, indent=2).encode("utf-8")

def _warning_codes_by_row(
    value: dict[str, Any],
    document_type: str,
) -> dict[int, set[str]]:
    codes_by_row: dict[int, set[str]] = {}

    if document_type == "holerite":
        pages = [
            PayrollPage.model_validate(page)
            for page in value.get("pages", [])
        ]

        warnings = validate_payroll_pages(pages)

        for warning in warnings:
            if warning.row is not None:
                codes_by_row.setdefault(
                    warning.row,
                    set(),
                ).add(warning.code)

        return codes_by_row

    if document_type == "cartao-ponto":
        warnings = validate_timecard_transcription(value)
        row_positions: dict[tuple[int, int], int] = {}
        export_row = 0

        for page in value.get("pages", []):
            page_number = page.get("page", 0)

            for page_row, _day in enumerate(
                page.get("days", [])
            ):
                row_positions[
                    (page_number, page_row)
                ] = export_row
                export_row += 1

        for warning in warnings:
            if warning.row is None:
                continue

            position = row_positions.get(
                (warning.page, warning.row)
            )

            if position is not None:
                codes_by_row.setdefault(
                    position,
                    set(),
                ).add(warning.code)

    return codes_by_row


def export_csv(value: dict[str, Any], document_type: str) -> bytes:
    headers, rows = _rows(value, document_type)
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(headers)
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8-sig")


def export_xlsx(
    value: dict[str, Any],
    document_type: str,
) -> bytes:
    headers, rows = _rows(value, document_type)
    warning_codes = _warning_codes_by_row(
        value,
        document_type,
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transcrição"
    sheet.append(headers)

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(
            color="FFFFFF",
            bold=True,
        )

    sequential_codes = {
        "NON_SEQUENTIAL_DATE",
        "NON_SEQUENTIAL_MONTH",
    }

    for row_index, values in enumerate(rows):
        sheet.append(values)
        row_number = sheet.max_row
        codes = warning_codes.get(row_index, set())

        has_uncertainty = any(
            "?" in str(item)
            for item in values
        )

        odd_punches = (
            document_type == "cartao-ponto"
            and len(
                [
                    item
                    for item in values[1:]
                    if item
                ]
            )
            % 2
            == 1
        )

        has_sequential_error = bool(
            codes & sequential_codes
        )

        has_warning = (
            has_uncertainty
            or odd_punches
            or bool(codes)
        )

        if has_sequential_error:
            for cell in sheet[row_number]:
                cell.fill = ERROR_FILL

            sheet.cell(
                row=row_number,
                column=1,
            ).border = ERROR_BORDER

        elif has_warning:
            for cell in sheet[row_number]:
                cell.fill = WARNING_FILL

    sheet.freeze_panes = "A2"

    for column in sheet.columns:
        width = min(
            max(
                len(str(cell.value or ""))
                for cell in column
            )
            + 2,
            45,
        )

        sheet.column_dimensions[
            column[0].column_letter
        ].width = width

    output = io.BytesIO()
    workbook.save(output)

    return output.getvalue()
